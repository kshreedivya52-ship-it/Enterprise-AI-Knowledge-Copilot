import os
import sys
import json
import logging
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Configure logging
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("seed_data")

# Add project root to sys.path
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.append(PROJECT_ROOT)

from app.core.chunker import chunk_document
from app.core.vector_db import upsert_document_chunks, init_qdrant


def escape_pdf_text(text: str) -> str:
    """Escapes special PDF string characters."""
    return text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')


def build_pdf(filepath: str, title: str, paragraphs: list):
    """
    Builds a standard, multi-page valid PDF 1.4 document from paragraphs.
    Handles text wrapping, page breaks, headings, and font switching.
    """
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    pages_data = []
    current_page_ops = []
    y = 740

    # Header title on first page
    current_page_ops.append(f'BT /F1 14 Tf 50 {y} Td ({escape_pdf_text(title[:72])}) Tj ET')
    y -= 26

    for p in paragraphs:
        p_clean = p.strip()
        if not p_clean:
            y -= 8
            continue

        is_heading = p_clean.startswith('##') or p_clean.startswith('SECTION') or p_clean.isupper()
        if is_heading:
            p_clean = p_clean.lstrip('#').strip()
            y -= 10
            font = '/F1 11'
            line_height = 15
        else:
            font = '/F2 9.5'
            line_height = 13

        # Word wrap text to ~82 characters per line
        words = p_clean.split()
        lines = []
        cur_line = []
        for w in words:
            if len(' '.join(cur_line + [w])) > 82:
                lines.append(' '.join(cur_line))
                cur_line = [w]
            else:
                cur_line.append(w)
        if cur_line:
            lines.append(' '.join(cur_line))

        for line in lines:
            if y < 55:
                # Add page footer before rolling over
                current_page_ops.append(f'BT /F2 8 Tf 50 35 Td (Nexus Enterprise Technologies Inc. - Internal Corporate Document) Tj ET')
                pages_data.append(current_page_ops)
                current_page_ops = []
                y = 740

            current_page_ops.append(f'BT {font} Tf 50 {y} Td ({escape_pdf_text(line)}) Tj ET')
            y -= line_height

        y -= 6

    # Close last page
    if current_page_ops:
        current_page_ops.append(f'BT /F2 8 Tf 50 35 Td (Nexus Enterprise Technologies Inc. - Internal Corporate Document) Tj ET')
        pages_data.append(current_page_ops)

    num_pages = len(pages_data)
    objects = []
    # 1: Catalog
    objects.append(b'<< /Type /Catalog /Pages 2 0 R >>')

    # Page and Content object IDs
    page_obj_ids = [3 + i * 2 for i in range(num_pages)]
    content_obj_ids = [4 + i * 2 for i in range(num_pages)]
    font1_id = 3 + num_pages * 2
    font2_id = font1_id + 1

    kids_str = ' '.join(f'{pid} 0 R' for pid in page_obj_ids)
    pages_obj = f'<< /Type /Pages /Kids [{kids_str}] /Count {num_pages} >>'.encode('latin1')
    objects.append(pages_obj)

    for i in range(num_pages):
        page_content = '\n'.join(pages_data[i]).encode('latin1', errors='replace')
        page_dict = (
            f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
            f'/Resources << /Font << /F1 {font1_id} 0 R /F2 {font2_id} 0 R >> >> '
            f'/Contents {content_obj_ids[i]} 0 R >>'
        ).encode('latin1')
        content_dict = f'<< /Length {len(page_content)} >>\nstream\n'.encode('latin1') + page_content + b'\nendstream'
        objects.append(page_dict)
        objects.append(content_dict)

    # Fonts: Helvetica-Bold (F1) and Helvetica (F2)
    objects.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>')
    objects.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')

    # Assemble PDF with cross-reference table and trailer
    out = [b'%PDF-1.4\n']
    offsets = []
    curr_offset = len(out[0])
    for idx, obj in enumerate(objects, start=1):
        offsets.append(curr_offset)
        chunk = f'{idx} 0 obj\n'.encode('latin1') + obj + b'\nendobj\n'
        out.append(chunk)
        curr_offset += len(chunk)

    xref_offset = curr_offset
    out.append(f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'.encode('latin1'))
    for off in offsets:
        out.append(f'{off:010d} 00000 n \n'.encode('latin1'))
    out.append(
        f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n'.encode('latin1')
    )

    with open(filepath, 'wb') as f:
        f.write(b''.join(out))


def build_docx(filepath: str, title: str, sections: list):
    """Generates a professional Word .docx document."""
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=0)
    for sec_title, sec_body in sections:
        doc.add_heading(sec_title, level=1)
        if isinstance(sec_body, list):
            for para in sec_body:
                doc.add_paragraph(para)
        else:
            doc.add_paragraph(sec_body)
    doc.save(filepath)


def build_xlsx(filepath: str, sheets_data: dict):
    """Generates a styled multi-tab Excel workbook."""
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    border_thin = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for sheet_name, rows in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name)
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                if r_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    if isinstance(val, (int, float)):
                        cell.alignment = Alignment(horizontal='right')
                        cell.number_format = '$#,##0' if val > 1000 else '#,##0'
                    else:
                        cell.alignment = Alignment(horizontal='left')
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(filepath)


def build_md(filepath: str, content: str):
    """Writes a standard Markdown file."""
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(content)


# ==============================================================================
# DOCUMENT DEFINITIONS & CONTENTS FOR NEXUS ENTERPRISE TECHNOLOGIES INC.
# ==============================================================================

DOCUMENTS = {
    # --------------------------------------------------------------------------
    # 1. HR AND EMPLOYEE DOCUMENTS (department="hr")
    # --------------------------------------------------------------------------
    "hr": [
        {
            "filename": "Employee_Handbook.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - HR Employee Handbook & Operational Guide",
            "text": """# Nexus Technologies - HR Employee Handbook & Operational Guide (2026 Edition)

## SECTION 1: ORGANIZATIONAL MISSION AND HR CHARTER
Nexus Enterprise Technologies Inc. is committed to fostering an equitable, high-performance workplace culture. The Human Resources division oversees talent acquisition, organizational effectiveness, employee total rewards, career pathing, and workplace conflict mediation.

## SECTION 2: PERFORMANCE MANAGEMENT AND APPRAISAL CYCLES
Performance reviews are conducted semi-annually across all divisions. The mid-year review cycle runs in June, and the end-of-year comprehensive review concludes in December. Employees are evaluated across a 5-point rating scale: 1 (Unsatisfactory), 2 (Needs Improvement), 3 (Consistently Meets Expectations), 4 (Exceeds Expectations), and 5 (Outstanding Impact). Merit-based compensation adjustments range from 3% to 12% annually based on individual and company performance tiers.

## SECTION 3: PERFORMANCE IMPROVEMENT PLAN (PIP) ESCALATION
When an employee's performance fails to meet role expectations, a formal Performance Improvement Plan (PIP) is initiated by the direct manager in consultation with the HR Business Partner (HRBP). The PIP duration is strictly 30, 45, or 60 calendar days with clear measurable milestones. Weekly review check-ins are mandatory. Failure to achieve the stipulated objectives by day 60 will result in formal employment separation.

## SECTION 4: EQUITY GRANTS AND VESTING SCHEDULE
Standard corporate equity grants (Restricted Stock Units - RSUs) are governed by a 4-year vesting schedule with a 1-year cliff. Twenty-five percent (25%) of the awarded shares vest upon completion of the first anniversary, followed by quarterly vesting of 6.25% for the remaining 36 months. Equity grants require continuous active service and are managed through Carta.

## SECTION 5: CONFIDENTIAL EMPLOYEE ASSISTANCE PROGRAM (EAP)
Nexus provides 24/7 confidential mental health support, family counseling, legal consultation, and financial coaching through our global EAP partner Lyra Health. All employees and their dependents receive up to 12 free individual therapy sessions per calendar year with complete confidentiality."""
        },
        {
            "filename": "HR_Policies_2026.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Internal HR Operational Policies (2026)",
            "text": """# Nexus Technologies - Internal HR Operational Policies (2026)

## SECTION 1: TALENT REQUISITION AND RECRUITING PROTOCOL
All headcount additions must be budgeted and approved through the HR Headcount Committee consisting of the VP of HR, Finance Controller, and Department VP. Standard job requisitions must be posted internally on Greenhouse for a minimum of 7 calendar days before external candidate sourcing begins. Background checks through Checkr are mandatory and verify 7-year criminal history, educational credentials, and past employment verification.

## SECTION 2: CORPORATE SALARY BENCHMARKING BANDS (2026)
Compensation bands are refreshed annually using Radford and Mercer salary surveys to target the 75th percentile of the technology market:
- Level 1 (Associate): $85,000 - $110,000 base salary + 10% bonus target.
- Level 2 (Mid-Level): $115,000 - $145,000 base salary + 15% bonus target.
- Level 3 (Senior): $150,000 - $185,000 base salary + 20% bonus target.
- Level 4 (Staff / Lead): $190,000 - $235,000 base salary + 25% bonus target.
- Level 5 (Principal / Manager): $240,000 - $290,000 base salary + 30% bonus target.
- Level 6 (Director / Executive): $300,000 - $380,000 base salary + 40% bonus target.

## SECTION 3: SEVERANCE AND EMPLOYMENT TERMINATION GUIDELINES
In the event of involuntary reduction in force (RIF) or mutual separation, eligible employees receive severance pay calculated as 2 weeks of base salary per completed year of tenure, with a minimum guaranteed floor of 4 weeks and a maximum ceiling of 26 weeks. Company-subsidized COBRA health coverage continues for the duration of the severance payment period.

## SECTION 4: WORKPLACE HARASSMENT AND GRIEVANCE INVESTIGATIONS
The HR Employee Relations committee handles all workplace grievances. Reports submitted via the confidential hotline (hr-grievance@nexustech.internal) must acknowledge receipt within 24 hours. Formal investigations must conclude within 10 business days, culminating in a written determination memo signed by the Chief People Officer."""
        },
        {
            "filename": "Leave_and_Attendance_Policy.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Leave and Attendance Operational Policy",
            "text": """# Nexus Technologies - Leave and Attendance Operational Policy

## SECTION 1: PAID TIME OFF (PTO) ACCRUAL AND ROLLOVER
Full-time salaried employees accrue Paid Time Off at a rate of 1.83 days per completed month, totaling 22 days per calendar year. Up to 10 unused PTO days may be rolled over into the subsequent calendar year, but rolled-over days must be utilized by March 31 of the new fiscal year or they are automatically forfeited. PTO requests of 3 consecutive days or more must be submitted in BambooHR at least two weeks in advance.

## SECTION 2: PAID SICK LEAVE AND WELLNESS DAYS
Employees receive 12 dedicated paid sick leave days per calendar year credited on January 1. Sick leave may be used for personal illness, medical appointments, or caring for an immediate family member. Medical documentation from a licensed healthcare provider is required for continuous sick leaves exceeding 3 business days.

## SECTION 3: PARENTAL AND CAREGIVER LEAVE
Nexus offers 16 weeks of 100% fully paid parental leave for primary caregivers following the birth, adoption, or foster placement of a child. Secondary caregivers receive 8 weeks of 100% fully paid leave. Parental leave may be taken continuously or in two separate blocks within the first 12 months of the child's arrival. Health and retirement benefits continue uninterrupted during the leave period.

## SECTION 4: COMPASSIONATE BEREAVEMENT AND JURY DUTY LEAVE
Employees are eligible for up to 5 consecutive paid days for the death of an immediate family member (spouse, child, parent, sibling) and 3 paid days for extended family members. Nexus provides up to 10 business days of paid leave for mandated civic jury service upon provision of official court documentation.

## SECTION 5: UNPAID SABBATICAL ELIGIBILITY
Employees who have completed 3 or more continuous years of service in good standing may apply for an unpaid career sabbatical ranging from 30 to 90 calendar days for educational pursuit, volunteering, or personal enrichment with VP approval."""
        },
        {
            "filename": "Employee_Onboarding_Guide.docx",
            "type": "docx",
            "title": "Nexus Technologies - New Hire 30-60-90 Day Onboarding Guide",
            "sections": [
                ("1. Day One Welcome & Provisioning", [
                    "Welcome to Nexus Enterprise Technologies Inc.! On Day 1, new team members will complete I-9 identity verification, receive their MacBook Pro or Lenovo ThinkPad workstation, configure their Okta Single Sign-On account, and activate corporate Slack and Gmail credentials.",
                    "An IT orientation session takes place at 10:00 AM CST to ensure YubiKey MFA enrollment and GlobalProtect VPN installation are configured."
                ]),
                ("2. Week One: Orientation & Compliance Certifications", [
                    "During Week 1, new hires will meet their assigned Onboarding Buddy, complete the mandatory HR Code of Conduct training on KnowBe4, review the departmental OKRs, and attend the executive welcome session with VP of HR Rachel Sterling and CEO Marcus Vance."
                ]),
                ("3. The 30-Day Milestone: Tooling and First Sprint Contribution", [
                    "By Day 30, the new hire should have full familiarity with the internal tech stack (GitHub Enterprise, Jira, Confluence, AWS dev accounts), complete their first sprint assignment or departmental deliverable, and conduct their initial 30-day alignment check-in with their direct manager."
                ]),
                ("4. The 60-Day Milestone: Independent Project Ownership", [
                    "By Day 60, employees are expected to operate with minimal supervision, manage individual tasks, participate actively in cross-functional standups, and provide preliminary feedback on the team onboarding experience."
                ]),
                ("5. The 90-Day Milestone: Formal Probationary Review", [
                    "At Day 90, the manager and HRBP conduct the official 90-Day Milestone Review. Upon successful evaluation, the employee graduates from the probationary period, becomes eligible for tuition reimbursement, and establishes annual goals in Lattice."
                ])
            ],
            "text": """# Nexus Technologies - New Hire 30-60-90 Day Onboarding Guide

## 1. Day One Welcome & Provisioning
Welcome to Nexus Enterprise Technologies Inc.! On Day 1, new team members will complete I-9 identity verification, receive their MacBook Pro or Lenovo ThinkPad workstation, configure their Okta Single Sign-On account, and activate corporate Slack and Gmail credentials. An IT orientation session takes place at 10:00 AM CST to ensure YubiKey MFA enrollment and GlobalProtect VPN installation are configured.

## 2. Week One: Orientation & Compliance Certifications
During Week 1, new hires will meet their assigned Onboarding Buddy, complete the mandatory HR Code of Conduct training on KnowBe4, review the departmental OKRs, and attend the executive welcome session with VP of HR Rachel Sterling and CEO Marcus Vance.

## 3. The 30-Day Milestone: Tooling and First Sprint Contribution
By Day 30, the new hire should have full familiarity with the internal tech stack (GitHub Enterprise, Jira, Confluence, AWS dev accounts), complete their first sprint assignment or departmental deliverable, and conduct their initial 30-day alignment check-in with their direct manager.

## 4. The 60-Day Milestone: Independent Project Ownership
By Day 60, employees are expected to operate with minimal supervision, manage individual tasks, participate actively in cross-functional standups, and provide preliminary feedback on the team onboarding experience.

## 5. The 90-Day Milestone: Formal Probationary Review
At Day 90, the manager and HRBP conduct the official 90-Day Milestone Review. Upon successful evaluation, the employee graduates from the probationary period, becomes eligible for tuition reimbursement, and establishes annual goals in Lattice."""
        },
        {
            "filename": "Work_From_Home_Policy.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Work From Home & Remote Work Policy (HR Department)",
            "text": """# Nexus Technologies - Work From Home & Remote Work Policy (HR Department)

## SECTION 1: HYBRID WORK MODEL FRAMEWORK
Nexus operates a 3-2 hybrid workplace policy. Employees assigned to physical office hubs (Austin, Seattle, London) work in-office three days a week (typically Monday, Wednesday, and Thursday) and may work remotely on Tuesdays and Fridays. Department directors may adjust team anchor days based on project requirements.

## SECTION 2: FULL-REMOTE EXCEPTION APPROVALS
Employees requesting 100% remote work status must have completed a minimum of 6 months of service with a performance rating of 'Exceeds Expectations' or higher. Full-remote transitions require written approval from the Department Vice President and Chief People Officer to assess payroll tax, legal nexus, and timezone collaboration viability.

## SECTION 3: HOME OFFICE ERGONOMIC STIPEND
To ensure a healthy working environment, all permanent employees receive a one-time tax-free home office ergonomic stipend of $750 upon hire. This stipend can be claimed via Expensify for purchasing ergonomic office chairs, standing desks, external monitors, or lighting accessories.

## SECTION 4: HIGH-SPEED INTERNET REIMBURSEMENT
Employees working hybrid or remote are eligible for an ongoing monthly broadband reimbursement of $75 to offset domestic high-speed internet expenses. Claims must be submitted bi-monthly with an attached telecommunications service provider invoice.

## SECTION 5: REMOTE WORKING FROM ABROAD
Employees in good standing may work outside their designated country of employment for up to 30 calendar days per fiscal year, subject to prior manager notice and HR compliance verification regarding international tax residence and export control laws."""
        }
    ],

    # --------------------------------------------------------------------------
    # 2. FINANCE AND ACCOUNTING DOCUMENTS (department="finance")
    # --------------------------------------------------------------------------
    "finance": [
        {
            "filename": "Annual_Financial_Report_2026.pdf",
            "type": "pdf",
            "title": "Nexus Technologies Inc. - FY2026 Annual Financial Report",
            "text": """# Nexus Technologies Inc. - FY2026 Annual Financial Report

## SECTION 1: FINANCIAL PERFORMANCE HIGHLIGHTS
Nexus Enterprise Technologies Inc. completed Fiscal Year 2026 with record performance:
- Total Revenue reached $142.8 million, representing a 24.2% YoY increase compared to FY2025 ($115.0 million).
- Annual Recurring Revenue (ARR) closed at $128.5 million with a Net Dollar Retention (NDR) rate of 118.4%.
- Non-GAAP Gross Profit reached $111.7 million, delivering a strong Gross Margin of 78.2%.
- Adjusted EBITDA finished at $28.4 million, establishing a healthy 19.9% EBITDA margin.

## SECTION 2: OPERATING EXPENSES ANALYSIS
Consolidated Operating Expenses for FY2026 totaled $83.3 million:
- Research & Development (R&D): $42.6 million (29.8% of revenue), reflecting continuous investment in scalable hybrid vector search and generative enterprise copilot features.
- Sales & Marketing (S&M): $26.8 million (18.8% of revenue), focused on enterprise Fortune 500 account expansion.
- General & Administrative (G&A): $13.9 million (9.7% of revenue), covering legal, accounting, compliance, and corporate facilities.

## SECTION 3: CASH POSITION AND BALANCE SHEET
As of December 31, 2026, Nexus held $64.2 million in unrestricted cash, cash equivalents, and short-term US Treasury securities. Free Cash Flow generated during the fiscal year was $22.7 million. Long-term debt obligations remained at $0.00, preserving our debt-free capital structure.

## SECTION 4: FY2027 FINANCIAL GUIDANCE
For Fiscal Year 2027, the executive leadership projects total revenue between $175.0 million and $182.0 million (23% - 27% growth), with non-GAAP operating margin expanding to 22.0% and targeted ARR reaching $160 million by Q4 2027."""
        },
        {
            "filename": "Expense_Reimbursement_Policy.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Corporate Expense Reimbursement Policy",
            "text": """# Nexus Technologies - Corporate Expense Reimbursement Policy

## SECTION 1: GENERAL PRINCIPLES AND FILING TIMELINES
All business-related expenses incurred on behalf of Nexus Technologies must be legitimate, reasonable, and necessary. Expense claims must be submitted via the corporate Brex/Expensify system within 30 calendar days from the date of expenditure. Claims submitted after 60 days will be rejected unless accompanied by executive CFO exemption approval.

## SECTION 2: RECEIPT THRESHOLDS AND DOCUMENTATION
Itemized receipts showing transaction date, vendor name, item description, and total amount paid are strictly required for any expenditure exceeding $25.00 USD. Credit card statements without itemized receipts will not be accepted. Alcoholic beverages must be itemized separately and are restricted to team dinners pre-approved by a Director.

## SECTION 3: LODGING AND GROUND TRANSPORTATION LIMITS
- Domestic Lodging: Standard hotel room cap is $225 per night (exclusive of local taxes). For high-cost metropolitan markets (San Francisco, New York City, London, Zurich, Tokyo), the allowable cap is $340 per night.
- Ground Transportation: Employees should utilize ride-sharing services (Uber/Lyft Standard) or public transit. Premium ride tiers (Uber Black, SUV) are non-reimbursable. Personal vehicle business mileage is reimbursed at the official IRS standard rate of $0.67 per mile.

## SECTION 4: AUDIT PROCEDURES AND FRAUD PREVENTION
The Internal Audit department randomly samples 15% of all monthly expense claims for compliance verification. Any deliberate submission of fraudulent, inflated, or duplicated expense claims constitutes gross misconduct resulting in immediate termination and potential legal prosecution."""
        },
        {
            "filename": "Invoice_Processing_Guidelines.docx",
            "type": "docx",
            "title": "Nexus Technologies - Accounts Payable Invoice Processing SOP",
            "sections": [
                ("1. Purpose & Scope", [
                    "This Standard Operating Procedure defines the guidelines for Accounts Payable (AP) staff to receive, validate, record, and disburse payments for all vendor and supplier invoices in NetSuite ERP."
                ]),
                ("2. The Three-Way Matching Protocol", [
                    "Every invoice submitted for payment must satisfy the Three-Way Matching rule before payment authorization:",
                    "1. Purchase Order (PO): Valid PO created in Coupa and signed according to Delegation of Authority.",
                    "2. Receiving Receipt: Formal goods receipt or service milestone sign-off by the business owner.",
                    "3. Vendor Invoice: Itemized billing document with matching line items, quantities, and agreed pricing within a 2% tolerance threshold."
                ]),
                ("3. Standard Payment Terms and Schedule", [
                    "Standard corporate payment terms are Net-45 days from receipt of a verified invoice, unless Net-30 has been formally executed in an approved Master Vendor Agreement. AP payment batches are executed every Thursday via automated ACH transfer or domestic Fedwire."
                ]),
                ("4. Tax Form Collection & Dispute Handling", [
                    "No payment may be disbursed to any vendor without a verified W-9 (for US entities) or W-8BEN / W-8BEN-E (for non-US entities) on file with the AP team. Disputed invoices must be placed on 'AP Hold' in NetSuite within 48 hours of receipt, and notice must be dispatched to the vendor contact."
                ])
            ],
            "text": """# Nexus Technologies - Accounts Payable Invoice Processing SOP

## 1. Purpose & Scope
This Standard Operating Procedure defines the guidelines for Accounts Payable (AP) staff to receive, validate, record, and disburse payments for all vendor and supplier invoices in NetSuite ERP.

## 2. The Three-Way Matching Protocol
Every invoice submitted for payment must satisfy the Three-Way Matching rule before payment authorization:
1. Purchase Order (PO): Valid PO created in Coupa and signed according to Delegation of Authority.
2. Receiving Receipt: Formal goods receipt or service milestone sign-off by the business owner.
3. Vendor Invoice: Itemized billing document with matching line items, quantities, and agreed pricing within a 2% tolerance threshold.

## 3. Standard Payment Terms and Schedule
Standard corporate payment terms are Net-45 days from receipt of a verified invoice, unless Net-30 has been formally executed in an approved Master Vendor Agreement. AP payment batches are executed every Thursday via automated ACH transfer or domestic Fedwire.

## 4. Tax Form Collection & Dispute Handling
No payment may be disbursed to any vendor without a verified W-9 (for US entities) or W-8BEN / W-8BEN-E (for non-US entities) on file with the AP team. Disputed invoices must be placed on 'AP Hold' in NetSuite within 48 hours of receipt, and notice must be dispatched to the vendor contact."""
        },
        {
            "filename": "Department_Budget_2026.xlsx",
            "type": "xlsx",
            "title": "Nexus Technologies - FY2026 Department Budgets",
            "sheets": {
                "Summary": [
                    ["Department Name", "Cost Center", "FY26 Budget ($)", "FY25 Actual ($)", "YoY Growth (%)", "Headcount (FTE)"],
                    ["Engineering & Cloud", "CC-101", 48500000, 39200000, 23.7, 185],
                    ["Sales & Business Dev", "CC-201", 34100000, 28400000, 20.1, 120],
                    ["Marketing & Growth", "CC-301", 22400000, 18100000, 23.8, 65],
                    ["General IT & Operations", "CC-401", 12500000, 10800000, 15.7, 42],
                    ["Human Resources & People", "CC-501", 6800000, 5600000, 21.4, 28],
                    ["Finance, Tax & Accounting", "CC-601", 5200000, 4400000, 18.2, 22],
                    ["Legal, Risk & Compliance", "CC-701", 4300000, 3600000, 19.4, 16],
                    ["Total Enterprise Budget", "TOTAL", 133800000, 110100000, 21.5, 478]
                ],
                "Engineering_Detail": [
                    ["Subcategory", "Q1 Allocation", "Q2 Allocation", "Q3 Allocation", "Q4 Allocation", "Full Year Total"],
                    ["AWS & Cloud Infrastructure", 4200000, 4400000, 4700000, 4900000, 18200000],
                    ["Engineering Headcount & Salaries", 6200000, 6400000, 6600000, 6900000, 26100000],
                    ["Developer Tooling & Licenses", 650000, 700000, 720000, 730000, 2800000],
                    ["External Security & Pen Testing", 300000, 350000, 400000, 350000, 1400000],
                    ["Engineering Total", 11350000, 11850000, 12420000, 12880000, 48500000]
                ],
                "Finance_Detail": [
                    ["Expense Line", "Annual Budget", "Notes / Vendor Partner"],
                    ["External Financial Audit", 680000, "Deloitte LLP annual PCAOB audit fees"],
                    ["Tax Planning & Compliance", 420000, "International transfer pricing and state taxes"],
                    ["Treasury & Banking Fees", 310000, "Silicon Valley Bank and JPMorgan treasury lines"],
                    ["Corporate D&O Insurance", 1200000, "Comprehensive executive and cyber risk underwriting"],
                    ["NetSuite & Coupa ERP", 950000, "Enterprise licenses and cloud hosting support"]
                ]
            },
            "text": """# Nexus Technologies - FY2026 Department Budgets
The FY2026 Total Enterprise Budget is $133,800,000 across 478 full-time equivalents (FTE).
Key departmental budget allocations:
1. Engineering & Cloud (Cost Center CC-101): $48,500,000 (185 FTE). AWS & Cloud Infrastructure: $18.2M, Engineering Salaries: $26.1M, Developer Tools: $2.8M, Pen Testing: $1.4M.
2. Sales & Business Dev (Cost Center CC-201): $34,100,000 (120 FTE).
3. Marketing & Growth (Cost Center CC-301): $22,400,000 (65 FTE).
4. General IT & Operations (Cost Center CC-401): $12,500,000 (42 FTE).
5. Human Resources & People (Cost Center CC-501): $6,800,000 (28 FTE).
6. Finance, Tax & Accounting (Cost Center CC-601): $5,200,000 (22 FTE). External Audit: $680k, Tax: $420k, Insurance: $1.2M, ERP Software: $950k.
7. Legal, Risk & Compliance (Cost Center CC-701): $4,300,000 (16 FTE)."""
        },
        {
            "filename": "Purchase_Approval_Policy.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Corporate Purchase Approval & Delegation Matrix",
            "text": """# Nexus Technologies - Corporate Purchase Approval & Delegation Matrix

## SECTION 1: PURPOSE AND APPLICABILITY
This policy sets forth the mandatory delegation of authority limits for entering into financial commitments, purchase orders, vendor agreements, and capital expenditures (CapEx) on behalf of Nexus Technologies.

## SECTION 2: DELEGATION OF AUTHORITY (DOA) THRESHOLDS
All procurement orders in Coupa must adhere to the following sign-off approval matrix:
- Tier 1: Up to $1,000 USD requires Direct Manager approval.
- Tier 2: $1,001 to $10,000 USD requires Department Director approval.
- Tier 3: $10,001 to $50,000 USD requires Department Vice President AND Finance Controller approval.
- Tier 4: $50,001 to $250,000 USD requires Chief Financial Officer (CFO) approval.
- Tier 5: Above $250,000 USD requires Chief Executive Officer (CEO) AND Board of Directors Audit Committee approval.

## SECTION 3: COMPETITIVE BIDDING AND RFP REQUIREMENTS
Any vendor contract or technology commitment exceeding $50,000 USD in aggregate annual value must undergo a competitive Request for Proposal (RFP) process with a minimum of three qualified vendor proposals. Sole-source procurement exceptions must be justified in writing and approved by the CFO.

## SECTION 4: SPLITTING PURCHASES PROHIBITION
Structuring or splitting purchase orders into smaller denominations to bypass approval thresholds is strictly prohibited and constitutes an ethical violation subject to disciplinary action."""
        }
    ],

    # --------------------------------------------------------------------------
    # 3. TECHNICAL AND ENGINEERING DOCUMENTS (department="engineering")
    # --------------------------------------------------------------------------
    "engineering": [
        {
            "filename": "System_Architecture.pdf",
            "type": "pdf",
            "title": "Nexus Platform - High-Level System Architecture Specification",
            "text": """# Nexus Platform - High-Level System Architecture Specification

## SECTION 1: ARCHITECTURE OVERVIEW & INGRESS TIER
The Nexus Cloud Platform is built on a distributed microservices architecture hosted across multiple Amazon Web Services (AWS) Availability Zones in us-west-2 (Oregon) with an active-standby disaster recovery presence in us-east-1. Ingress traffic is fronted by Cloudflare Enterprise for DDoS mitigation and SSL termination, flowing through an Envoy API Gateway cluster with mTLS service mesh routing.

## SECTION 2: EVENT-DRIVEN MESSAGE BUS & SERVICES TIER
Core services communicate asynchronously via an Apache Kafka 3.6 event bus deployed with 3 brokers and a replication factor of 3. Core microservices include:
- Authentication & RBAC Service (FastAPI / Auth0 / Local JWT RS256/HS256)
- Ingestion Pipeline Engine (LlamaParse, async chunking, metadata enrichment)
- Vector Search & Reranking Service (Qdrant hybrid dense-sparse, BGE-Reranker-Large)
- LLM Grounding & Citation Synthesizer (LangChain, OpenAI GPT-4o, Google Gemini Flash)
- Chat History Persistence Service (PostgreSQL 16 with async SQLAlchemy)

## SECTION 3: CACHING & PERSISTENCE TIERS
- Primary Relational Database: PostgreSQL 16 hosted on AWS Aurora with Citus multi-node distributed tables and monthly partitioning on chat history.
- Caching Layer: Redis Cluster 7.2 configured with write-through caching, 4-hour TTL for session state, and LRU eviction policy.
- Vector Engine: Qdrant cluster storing 384-dimensional dense vectors (BAAI/bge-small-en-v1.5) and SPLADE sparse lexical vectors, enabling Reciprocal Rank Fusion (RRF) search.

## SECTION 4: OBSERVABILITY & DISASTER RECOVERY
Observability is maintained through OpenTelemetry instrumentation streaming to Datadog and Prometheus/Grafana. P99 API response latency is maintained below 320ms. The platform satisfies a Recovery Point Objective (RPO) of <1 minute and a Recovery Time Objective (RTO) of <15 minutes."""
        },
        {
            "filename": "API_Documentation.docx",
            "type": "docx",
            "title": "Nexus Platform - Core REST and gRPC API Reference Manual",
            "sections": [
                ("1. Authentication & Security", [
                    "All API requests must transmit an HTTP Authorization Bearer token header obtained from the authentication endpoint. Tokens are signed JWTs containing user claims (`sub`, `role`, `department`).",
                    "Example: Authorization: Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9..."
                ]),
                ("2. Ingestion Endpoint: POST /ingest/upload", [
                    "Allows authorized editors to upload corporate documents for parsing and vector indexing.",
                    "Request: Multipart form-data with `file` (binary) and optional `department` string (e.g. 'engineering', 'finance', 'hr', 'legal', 'general').",
                    "Response (200 OK): `{'message': 'Document ingested', 'filename': 'guide.pdf', 'department': 'engineering', 'total_chunks': 14}`"
                ]),
                ("3. Hybrid Search Endpoint: GET /search", [
                    "Performs hybrid vector retrieval, cross-encoder reranking, and grounded LLM answer generation.",
                    "Parameters: `q` (string, required): Search query. `session_id` (UUID, optional): Chat session ID for contextual conversational history.",
                    "Response (200 OK): `{'session_id': 'uuid', 'answer': '...', 'has_sufficient_context': true, 'citations': [{'source_file': '...', 'quote': '...'}]}`"
                ]),
                ("4. Rate Limiting and HTTP Error Codes", [
                    "Rate limits are enforced at the API Gateway: Standard tier permits 1,000 requests/minute; Enterprise tier permits 10,000 requests/minute.",
                    "Common Status Codes: 200 OK, 400 Bad Request, 401 Unauthorized, 403 Forbidden (RBAC violation), 429 Too Many Requests, 500 Internal Server Error."
                ])
            ],
            "text": """# Nexus Platform - Core REST and gRPC API Reference Manual

## 1. Authentication & Security
All API requests must transmit an HTTP Authorization Bearer token header obtained from the authentication endpoint. Tokens are signed JWTs containing user claims (`sub`, `role`, `department`). Example: Authorization: Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9...

## 2. Ingestion Endpoint: POST /ingest/upload
Allows authorized editors to upload corporate documents for parsing and vector indexing.
Request: Multipart form-data with `file` (binary) and optional `department` string (e.g. 'engineering', 'finance', 'hr', 'legal', 'general').
Response (200 OK): `{'message': 'Document ingested', 'filename': 'guide.pdf', 'department': 'engineering', 'total_chunks': 14}`

## 3. Hybrid Search Endpoint: GET /search
Performs hybrid vector retrieval, cross-encoder reranking, and grounded LLM answer generation.
Parameters: `q` (string, required): Search query. `session_id` (UUID, optional): Chat session ID for contextual conversational history.
Response (200 OK): `{'session_id': 'uuid', 'answer': '...', 'has_sufficient_context': true, 'citations': [{'source_file': '...', 'quote': '...'}]}`

## 4. Rate Limiting and HTTP Error Codes
Rate limits are enforced at the API Gateway: Standard tier permits 1,000 requests/minute; Enterprise tier permits 10,000 requests/minute.
Common Status Codes: 200 OK, 400 Bad Request, 401 Unauthorized, 403 Forbidden (RBAC violation), 429 Too Many Requests, 500 Internal Server Error."""
        },
        {
            "filename": "Database_Design.pdf",
            "type": "pdf",
            "title": "Nexus Platform - Database Design and Data Modeling Blueprint",
            "text": """# Nexus Platform - Database Design and Data Modeling Blueprint

## SECTION 1: RELATIONAL DATABASE SCHEMA (POSTGRESQL 16)
The primary transactional database model includes the following core relational schemas:
- `users`: Stores user identity, hashed credentials, role (`admin`, `editor`, `viewer`), and assigned department (`general`, `engineering`, `finance`, `hr`, `legal`).
- `conversations`: Primary chat threads identified by UUIDv7, storing `user_id`, `title`, `created_at`, and `updated_at`.
- `chat_messages`: Stores message records with `conversation_id`, `role` (`user`, `assistant`), `content`, and JSONB `citations` metadata.
- `audit_logs`: Partitioned monthly by timestamp to record all access attempts, query text, and security authorization checks.

## SECTION 2: VECTOR DATABASE SCHEMA (QDRANT)
The Qdrant collection `hybrid_knowledge_base` is structured with dual vector spaces:
- Dense Vector Space: 384 dimensions using Cosine distance, generated by `BAAI/bge-small-en-v1.5` for semantic embeddings.
- Sparse Vector Space: SPLADE lexical weights (`prithivida/Splade_PP_en_v1`) for exact term and keyword recall.
- Payload Keyword Indices: Dedicated payload indices on `department` (enabling sub-millisecond RBAC filtering) and `source_file`.

## SECTION 3: REPLICATION, BACKUP, AND DISASTER RECOVERY
The PostgreSQL database runs with a Multi-AZ standby replica with synchronous replication. Point-in-time recovery (PITR) is maintained through continuous Write-Ahead Log (WAL) archiving to encrypted Amazon S3 buckets. Automated full snapshots execute daily at 02:00 UTC and are retained for 35 days in compliance with enterprise audit criteria."""
        },
        {
            "filename": "Software_Deployment_Guide.md",
            "type": "md",
            "title": "Nexus Engineering - Production CI/CD and Kubernetes Deployment Guide",
            "text": """# Nexus Engineering - Production CI/CD and Kubernetes Deployment Guide

## 1. GitOps Architecture & CI/CD Pipeline
Nexus uses a GitOps deployment strategy driven by GitHub Actions and ArgoCD. Every merge to the `main` branch triggers:
1. Automated unit and integration testing suite execution.
2. Static code analysis and dependency vulnerability scans (Trivy and SonarQube).
3. Multi-stage Docker image build pushed to AWS Elastic Container Registry (ECR) tagged with Git SHA.
4. ArgoCD syncs the updated image tag to the Kubernetes production cluster.

## 2. Kubernetes Cluster Topology (EKS 1.30)
Production microservices run on Amazon Elastic Kubernetes Service (EKS) across three Availability Zones (`us-west-2a`, `us-west-2b`, `us-west-2c`). Pods are governed by:
- Horizontal Pod Autoscaler (HPA) targeting 70% CPU and 80% memory utilization.
- Resource Requests: 500m CPU, 1Gi RAM; Limits: 2000m CPU, 4Gi RAM per pod.
- PodDisruptionBudgets (PDB) guaranteeing a minimum of 66% available pods during maintenance.

## 3. Canary Deployment Strategy
Deployments utilize progressive canary rollouts via Argo Rollouts:
- Phase 1: 10% traffic routed to canary pods for 15 minutes.
- Automated Analysis: Datadog metrics monitor HTTP 5xx error rates and P95 latency.
- Rollback Trigger: If HTTP 5xx error rate exceeds 0.5% or latency exceeds 450ms, the rollout is automatically aborted and traffic reverts 100% to stable pods.
- Phase 2: Traffic scales to 50% for 15 minutes, then promotes to 100%.

## 4. Secrets Management
Zero secrets are stored in Git. Application secrets are stored in HashiCorp Vault and injected into pods as ephemeral volumes using the Kubernetes External Secrets Operator (ESO)."""
        },
        {
            "filename": "Troubleshooting_Manual.pdf",
            "type": "pdf",
            "title": "Nexus Platform - Incident Response & Troubleshooting Manual",
            "text": """# Nexus Platform - Incident Response & Troubleshooting Manual

## SECTION 1: INCIDENT SEVERITY CLASSIFICATION
- SEV-1 (Critical Outage): Core service down, data corruption risk, or outage affecting >5% of active users. Target response time: <10 minutes.
- SEV-2 (Major Degradation): Degraded search performance, ingestion delays >15 minutes, or isolated regional latency. Target response time: <30 minutes.
- SEV-3 (Minor Defect): Non-critical bugs with available workarounds. Target response time: <4 hours.

## SECTION 2: HIGH DATABASE CONNECTION SATURATION RUNBOOK
Symptoms: API endpoints return HTTP 500 or `asyncpg.exceptions.TooManyConnectionsError`.
Diagnostic Steps:
1. Check active pgBouncer connection count: `SHOW POOLS;`
2. Inspect long-running blocking queries in PostgreSQL:
   `SELECT pid, now() - query_start AS duration, query FROM pg_stat_activity WHERE state != 'idle' ORDER BY duration DESC;`
3. Terminate stuck rogue query: `SELECT pg_terminate_backend(pid);`
4. If saturation persists, scale pgBouncer max client connections from 100 to 250 in Helm values.

## SECTION 3: QDRANT VECTOR SEARCH LATENCY SPIKE RUNBOOK
Symptoms: Hybrid search endpoint latency exceeds 1.5 seconds.
Diagnostic Steps:
1. Inspect Qdrant CPU and memory consumption via Grafana dashboard.
2. Verify if HNSW background indexing is saturating CPU threads.
3. Check memory swap utilization. If memory exceeds 85%, adjust `hnsw_config.ef_construct` or scale cluster node size.
4. Execute collection optimization trigger: `POST /collections/hybrid_knowledge_base/update_params` to compact segments."""
        }
    ],

    # --------------------------------------------------------------------------
    # 4. GENERAL DOCUMENTS EMPLOYEE GUIDELINES (department="general")
    # --------------------------------------------------------------------------
    "general": [
        {
            "filename": "Employee_Handbook.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - General Employee Handbook (Company-Wide)",
            "text": """# Nexus Technologies - General Employee Handbook (Company-Wide)

## SECTION 1: WELCOME FROM THE CEO
Welcome to Nexus Enterprise Technologies Inc.! Founded in 2018, our mission is to empower global knowledge workers with intelligent, secure enterprise copilot platforms. Our corporate culture rests upon three core pillars: Customer Obsession, Radical Transparency, and Extreme Ownership.

## SECTION 2: GLOBAL WORKPLACES AND FACILITIES
Nexus maintains corporate offices in Austin (Global Headquarters), Seattle, London, and Bangalore. Standard office access is managed via digital keycards and Okta Mobile badges. Guests and visitors must sign in via Envoy at the reception desk and wear a visitor badge while on premises.

## SECTION 3: EMPLOYEE WELLNESS AND LEARNING BENEFITS
Nexus provides all full-time employees with comprehensive perks:
- Annual Wellness Stipend: $500 per year for gym memberships, fitness trackers, or ergonomic gear.
- Learning & Development Grant: $1,200 annually for professional certifications, industry conferences, and books.
- Commuter Benefits: Pre-tax transit pass deductions and monthly public transit subsidies up to $100.

## SECTION 4: DRESS CODE AND WORKPLACE ENVIRONMENT
Nexus observes a 'smart casual' dress policy. Employees are encouraged to dress comfortably while maintaining a professional appearance appropriate for client meetings. We maintain a fragrance-free, inclusive environment across all office facilities."""
        },
        {
            "filename": "Code_of_Conduct.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Global Code of Business Conduct & Ethics",
            "text": """# Nexus Technologies - Global Code of Business Conduct & Ethics

## SECTION 1: COMMITMENT TO MUTUAL RESPECT AND DIVERSITY
Nexus is dedicated to maintaining a collaborative, harassment-free workplace. We do not tolerate discrimination or harassment based on race, color, gender, sexual orientation, religion, disability, or veteran status. Any form of bullying, intimidation, or retaliation will result in immediate disciplinary action up to termination.

## SECTION 2: CONFLICT OF INTEREST DISCLOSURES
Employees must avoid situations where personal interests conflict with their obligations to Nexus. Outside employment, board advisory roles, or consulting engagements require prior written disclosure and clearance from the Legal and Compliance Department. Personal investments exceeding 1% in direct competitors or suppliers are prohibited without General Counsel approval.

## SECTION 3: CONFIDENTIAL INFORMATION AND INTELLECTUAL PROPERTY
All software source code, customer records, technical specifications, and strategic roadmaps developed during employment belong exclusively to Nexus Technologies. Employees must safeguard proprietary information and never disclose internal data to third parties.

## SECTION 4: RESPONSIBLE USE OF AI TECHNOLOGIES
Employees utilizing generative AI tools must adhere to company security guardrails. Uploading sensitive company source code, confidential financials, or customer Personally Identifiable Information (PII) into public, unvetted AI models is strictly prohibited. Only enterprise-approved AI copilot instances may be used for corporate tasks."""
        },
        {
            "filename": "Work_From_Home_Policy.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - General Remote Work & Workplace Guidelines",
            "text": """# Nexus Technologies - General Remote Work & Workplace Guidelines

## SECTION 1: GENERAL FLEXIBLE WORK PRINCIPLES
Nexus supports flexible work arrangements that balance individual autonomy with team collaboration. All eligible employees may participate in our hybrid working rhythm, coordinating with team leaders to establish collaborative anchor days.

## SECTION 2: WORKSPACE SETUP AND ERGONOMIC HEALTH
Remote employees are expected to maintain a dedicated, distraction-free workspace. Workstations should feature adequate lighting, stable high-speed broadband connection (>50 Mbps download / 15 Mbps upload), and an ergonomically sound chair and desk arrangement. Employees should take regular 5-minute screen breaks every hour.

## SECTION 3: COMMUNICATION STANDARDS AND CORE HOURS
To facilitate seamless collaboration across distributed teams, core working hours are established between 10:00 AM and 4:00 PM in the employee's local time zone. Slack status should be kept current, and calendar availability should be maintained in Google Calendar.

## SECTION 4: WORKSTATION SECURITY ESSENTIALS
Corporate laptops must remain secured at all times. Screens must be locked immediately whenever stepping away from the device. Under no circumstances may family members or guests access corporate-issued hardware."""
        },
        {
            "filename": "Office_Hours_and_Holidays.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - 2026 Office Hours and Global Holiday Calendar",
            "text": """# Nexus Technologies - 2026 Office Hours and Global Holiday Calendar

## SECTION 1: STANDARD OFFICE OPERATING HOURS
Corporate office facilities are open to employees Monday through Friday from 8:00 AM to 6:00 PM local time. Standard business operating hours are 8:30 AM to 5:30 PM. Facilities are equipped with badge access 24/7 for authorized personnel requiring emergency maintenance access.

## SECTION 2: OFFICIAL 2026 PAID COMPANY HOLIDAYS (US CALENDAR)
Nexus observes the following 11 official paid holidays in 2026:
- New Year's Day: Thursday, January 1, 2026
- Martin Luther King Jr. Day: Monday, January 19, 2026
- Presidents' Day: Monday, February 16, 2026
- Memorial Day: Monday, May 25, 2026
- Juneteenth National Independence Day: Friday, June 19, 2026
- Independence Day Observed: Friday, July 3, 2026
- Labor Day: Monday, September 7, 2026
- Thanksgiving Day: Thursday, November 26, 2026
- Day After Thanksgiving: Friday, November 27, 2026
- Christmas Eve: Thursday, December 24, 2026
- Christmas Day: Friday, December 25, 2026

## SECTION 3: FLOATING CULTURAL HOLIDAYS
In addition to the official holiday calendar, all full-time employees receive two (2) Personal Floating Holidays per calendar year. Floating holidays can be utilized for cultural, religious, or personal observance with one week prior manager notice."""
        },
        {
            "filename": "Travel_and_Expense_Guidelines.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Employee Travel and Meal Guidelines",
            "text": """# Nexus Technologies - Employee Travel and Meal Guidelines

## SECTION 1: CORPORATE TRAVEL RESERVATIONS
All business travel including airline tickets, hotel accommodations, and rental vehicles must be booked through the corporate travel management portal (Navan / TripActions). Bookings must be made at least 14 days in advance of domestic travel to secure cost-effective pricing.

## SECTION 2: AIR TRAVEL POLICY
Economy class bookings are standard for all domestic and international flights with duration under 6 hours. For intercontinental or transatlantic flights exceeding 6 continuous hours in the air, employees are authorized to book Premium Economy. Business class requires written pre-approval from the CFO.

## SECTION 3: DAILY MEAL PER-DIEM ALLOWANCES
When traveling on official business, employees receive a daily meal allowance of up to $75.00 USD (or local currency equivalent):
- Breakfast: $15.00
- Lunch: $25.00
- Dinner: $35.00
Allowances may be balanced across the day (e.g. $10 breakfast and $40 dinner). Itemized receipts are required for all meals."""
        },
        {
            "filename": "IT_Helpdesk_Guide.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Employee IT Helpdesk & Self-Service Manual",
            "text": """# Nexus Technologies - Employee IT Helpdesk & Self-Service Manual

## SECTION 1: HOW TO CONTACT THE IT HELPDESK
The Nexus IT Helpdesk provides technical support Monday through Friday, 24 hours a day:
- Slack Channel: Submit requests in `#help-it` for real-time triage.
- IT Service Portal: Log tickets at `https://helpdesk.nexustech.internal`.
- Critical Emergency Hotline: Call 1-888-555-NXIT for urgent outages.

## SECTION 2: OKTA SINGLE SIGN-ON & PASSWORD MANAGEMENT
- Passwords must be at least 14 characters, containing uppercase, lowercase, numbers, and symbols.
- Passwords expire every 180 days. Self-service password resets are available at `https://auth.nexustech.internal/reset`.
- Multi-Factor Authentication (MFA) is strictly enforced via YubiKey hardware keys or the Okta Verify mobile application.

## SECTION 3: CONNECTING TO GLOBALPROTECT CORPORATE VPN
Access to internal databases, staging environments, and code repositories requires connection to the GlobalProtect VPN:
1. Launch GlobalProtect from the system menu bar.
2. Enter the portal address: `vpn.nexustech.internal`.
3. Complete Okta biometric or push notification authentication.
4. Verify connected status (blue shield icon).

## SECTION 4: REPORTING LOST OR STOLEN HARDWARE
If a corporate laptop, smartphone, or security token is lost or stolen, you must report it to IT Security (`security@nexustech.internal`) and your manager within 2 hours. The IT security team will immediately execute a remote MDM device wipe to prevent data compromise."""
        }
    ],

    # --------------------------------------------------------------------------
    # 5. LEGAL DOCUMENTS: CONTRACTS & COMPLIANCE (department="legal")
    # --------------------------------------------------------------------------
    "legal": [
        {
            "filename": "Master_Service_Agreement.pdf",
            "type": "pdf",
            "title": "Nexus Technologies Inc. - Standard Enterprise Master Service Agreement",
            "text": """# Nexus Technologies Inc. - Standard Enterprise Master Service Agreement (MSA)

## SECTION 1: PROVISION OF SERVICES AND SLA COMMITMENTS
Nexus Enterprise Technologies Inc. ('Provider') grants Customer a non-exclusive, non-transferable subscription right to access and use the Nexus Cloud Platform. Provider commits to delivering a Service Level Agreement (SLA) uptime of ninety-nine and nine-tenths percent (99.9%) measured on a monthly calendar basis, excluding scheduled maintenance windows. In the event uptime falls below 99.9%, Customer is entitled to proportionate service fee credits.

## SECTION 2: INTELLECTUAL PROPERTY RIGHTS & CUSTOMER DATA
Customer retains all right, title, and ownership interest in Customer Data uploaded to the Platform. Provider retains all intellectual property rights in the Platform, including software code, API endpoints, machine learning models, and documentation. Provider shall not use Customer Data to train foundation models without express written authorization.

## SECTION 3: MUTUAL INDEMNIFICATION
Provider agrees to defend, indemnify, and hold harmless Customer against any third-party claims alleging that Customer's authorized use of the Platform infringes any valid patent, copyright, or trademark. Customer agrees to indemnify Provider against claims arising from unauthorized data uploaded in breach of applicable laws.

## SECTION 4: LIMITATION OF LIABILITY AND GOVERNING LAW
Neither party shall be liable for indirect, incidental, punitive, or consequential damages. Each party's maximum aggregate liability arising under this Agreement shall be capped at the total subscription fees paid by Customer in the twelve (12) months preceding the incident. This Agreement is governed by the laws of the State of Delaware, with dispute resolution via binding AAA arbitration in Wilmington, Delaware."""
        },
        {
            "filename": "Non_Disclosure_Agreement_NDA.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Mutual Non-Disclosure Agreement (NDA)",
            "text": """# Nexus Technologies - Mutual Non-Disclosure Agreement (NDA)

## SECTION 1: DEFINITION OF CONFIDENTIAL INFORMATION
'Confidential Information' refers to any proprietary, non-public technical, commercial, financial, or strategic information disclosed by one party ('Disclosing Party') to the other ('Receiving Party'), whether orally, visually, or in tangible or electronic form, that is designated as confidential or should reasonably be understood to be confidential.

## SECTION 2: EXCLUSIONS FROM CONFIDENTIALITY
Confidential Information does not include information that:
1. Is or becomes publicly known through no breach of this Agreement.
2. Was already in the lawful possession of Receiving Party prior to disclosure.
3. Is independently developed by Receiving Party without reference to Confidential Information.
4. Is rightfully received from a third party without duty of confidentiality.

## SECTION 3: OBLIGATIONS AND STANDARD OF CARE
Receiving Party shall hold Confidential Information in strict confidence, exercising at least the same degree of care it uses to protect its own confidential materials, but no less than reasonable commercial diligence. Receiving Party may disclose Confidential Information only to employees, contractors, and legal advisors with a strict need to know.

## SECTION 4: TERM OF OBLIGATION
The confidentiality obligations under this Agreement remain in effect for a period of two (2) years following disclosure, provided that trade secrets and proprietary source code shall remain protected in perpetuity until lawful public disclosure."""
        },
        {
            "filename": "Vendor_Agreement.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Master Vendor and Supplier Services Agreement",
            "text": """# Nexus Technologies - Master Vendor and Supplier Services Agreement

## SECTION 1: ENGAGEMENT AND DELIVERABLES ACCEPTANCE
Vendor agrees to provide services and deliverables described in applicable Statements of Work (SOW). Nexus shall have ten (10) business days from delivery to evaluate and test deliverables. If deliverables fail to meet agreed specifications, Nexus may reject them and Vendor must remediate within five (5) business days at no additional cost.

## SECTION 2: SUPPLIER SECURITY & COMPLIANCE MANDATE
Vendor warrants that it maintains rigorous information security controls conforming to SOC 2 Type II or ISO/IEC 27001 standards. If Vendor processes Nexus data, it must encrypt data in transit (TLS 1.3) and at rest (AES-256) and immediately notify Nexus within 24 hours of any suspected security incident.

## SECTION 3: PAYMENT TERMS AND INVOICING
Vendor invoices must reference an active Nexus Purchase Order (PO). Invoices are payable under Net-45 terms following formal acceptance of deliverables. Invoices without valid PO numbers will be rejected.

## SECTION 4: INSURANCE REQUIREMENTS
Vendor must maintain Commercial General Liability insurance of at least $2,000,000 per occurrence, Workers' Compensation in compliance with statutory minimums, and Cyber Liability insurance of at least $5,000,000 per occurrence naming Nexus as an additional insured."""
        },
        {
            "filename": "Client_Service_Contract.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Client Professional Services Agreement",
            "text": """# Nexus Technologies - Client Professional Services Agreement

## SECTION 1: STATEMENTS OF WORK (SOW) & PROFESSIONAL SERVICES
Nexus will provide technical implementation, custom integration, and training services as described in executed Statements of Work (SOW). Each SOW specifies scope, timeline, milestones, personnel assignments, and deliverable acceptance criteria.

## SECTION 2: BILLING RATES AND EXPENSES
Unless a fixed-fee milestone schedule is specified in an SOW, professional services are billed on a Time & Materials (T&M) basis at the standard enterprise blended rate of $240.00 USD per hour. Reimbursable client travel expenses must be pre-approved in writing by the Client Project Sponsor.

## SECTION 3: CLIENT RESPONSIBILITIES & PROJECT DEPENDENCIES
Timely delivery depends upon Client providing timely access to technical facilities, test data, and technical personnel. Delays caused by Client dependencies exceeding 5 business days will result in an equitable timeline adjustment and potential remobilization fees.

## SECTION 4: CHANGE MANAGEMENT PROCEDURE
Any modification to project scope or deliverables must be documented in a written Change Order signed by both authorized project managers prior to initiating additional work."""
        },
        {
            "filename": "Partnership_Agreement.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Strategic Alliance and Reseller Agreement",
            "text": """# Nexus Technologies - Strategic Alliance and Reseller Agreement

## SECTION 1: APPOINTMENT AND PARTNER TIERS
Nexus appoints Partner as a non-exclusive reseller and implementation partner. Partner levels are categorized as:
- Silver Partner: Minimum $250k annual contract value (ACV); 15% reseller margin.
- Gold Partner: Minimum $1.0M annual contract value; 22% reseller margin + dedicated partner manager.
- Platinum Partner: Minimum $3.0M annual contract value; 30% reseller margin + joint executive sponsorship.

## SECTION 2: CO-MARKETING AND MARKET DEVELOPMENT FUNDS (MDF)
Gold and Platinum partners qualify for Market Development Funds (MDF) equal to up to 2.5% of generated ACV, usable for joint webinars, trade shows, and demand generation campaigns with Nexus approval.

## SECTION 3: TRADEMARK USAGE AND BRAND GUIDELINES
Partner receives a limited, revocable license to display Nexus logos and partner badges in marketing materials, strictly adhering to the Nexus Global Brand Identity Guidelines.

## SECTION 4: TERM AND TERMINATION
This Agreement has an initial term of one (1) year, renewing automatically unless either party provides sixty (60) days written notice of non-renewal. Either party may terminate for material breach with 30 days notice to cure."""
        },
        {
            "filename": "Legal_Compliance_Policy.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Corporate Legal Compliance Policy",
            "text": """# Nexus Technologies - Corporate Legal Compliance Policy

## SECTION 1: COMPLIANCE OVERSIGHT STRUCTURE
The Legal Compliance Committee, chaired by the General Counsel, oversees enterprise adherence to all domestic and international legal obligations. The Committee reports quarterly to the Board of Directors Audit Committee on regulatory posture, ongoing audits, and compliance risk assessments.

## SECTION 2: ANTITRUST AND FAIR COMPETITION
Nexus strictly complies with the Sherman Act, Clayton Act, and international competition laws. Employees must never discuss pricing, customer allocation, market division, or employee compensation with competitors. Attending industry conferences requires strict adherence to published agendas.

## SECTION 3: SANCTIONS SCREENING & EXPORT CONTROLS
Nexus products, APIs, and cloud services are subject to US Export Administration Regulations (EAR) and Department of Treasury Office of Foreign Assets Control (OFAC) sanctions. Customers, vendors, and partners must undergo automated screening against SDN and consolidated denied-parties lists prior to transaction authorization.

## SECTION 4: WHISTLEBLOWER PROTECTION & NON-RETALIATION
Employees may report legal or compliance violations anonymously through the EthicsPoint portal or hotline (1-800-461-9330). Nexus enforces an absolute zero-tolerance policy against retaliation against any individual who raises a concern in good faith."""
        },
        {
            "filename": "Data_Protection_and_Privacy_Policy.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Global Data Protection & Privacy Policy",
            "text": """# Nexus Technologies - Global Data Protection & Privacy Policy

## SECTION 1: REGULATORY ADHERENCE (GDPR & CCPA/CPRA)
Nexus maintains strict compliance with the European Union General Data Protection Regulation (GDPR), UK GDPR, and the California Consumer Privacy Act as amended by the CPRA. All personal data collected from users is processed transparently, lawfully, and solely for specified business purposes.

## SECTION 2: DATA SUBJECT RIGHTS AND REQUESTS (DSAR)
Individuals have the right to access, rectify, port, and request erasure of their personal data. All Data Subject Access Requests (DSAR) received at `privacy@nexustech.internal` must be verified and fulfilled within thirty (30) calendar days without charge.

## SECTION 3: CROSS-BORDER DATA TRANSFERS
Transfers of personal data originating in the European Economic Area (EEA) to destinations outside the EEA rely upon Standard Contractual Clauses (SCCs) approved by the European Commission, supplemented by transfer impact assessments.

## SECTION 4: DATA BREACH NOTIFICATION PROTOCOL
In the event of a confirmed personal data breach, Nexus will notify competent supervisory authorities (including lead DPAs) without undue delay and, where feasible, not later than seventy-two (72) hours after becoming aware of the incident, in strict accordance with GDPR Article 33."""
        },
        {
            "filename": "Regulatory_Compliance_Manual.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Cloud Infrastructure Regulatory Compliance Manual",
            "text": """# Nexus Technologies - Cloud Infrastructure Regulatory Compliance Manual

## SECTION 1: SOC 1 AND SOC 2 TYPE II ATTESTATION
Nexus completes annual independent SOC 1 and SOC 2 Type II examinations conducted by a licensed CPA firm. Attestation reports evaluate system controls against AICPA Trust Services Criteria: Security, Availability, Processing Integrity, Confidentiality, and Privacy over a minimum 6-month testing window.

## SECTION 2: ISO/IEC 27001:2022 CERTIFICATION
The Information Security Management System (ISMS) governing Nexus infrastructure is certified under ISO/IEC 27001:2022. Security controls encompass cryptography, access management, vulnerability patching, physical perimeter controls, and incident management.

## SECTION 3: HIPAA COMPLIANCE FOR HEALTHCARE DATA
When processing Protected Health Information (PHI) on behalf of covered entities, Nexus executes formal Business Associate Agreements (BAAs). PHI is protected through dedicated VPC isolation, AES-256 encryption at rest, TLS 1.3 in transit, and immutable audit logging.

## SECTION 4: PCI-DSS PAYMENT SECURITY
Nexus utilizes PCI-DSS Level 1 certified payment processors (Stripe) for credit card transactions, ensuring zero cardholder data touches or resides on Nexus application servers."""
        },
        {
            "filename": "Records_Retention_Policy.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Corporate Records Retention and Destruction Policy",
            "text": """# Nexus Technologies - Corporate Records Retention and Destruction Policy

## SECTION 1: PURPOSE AND CLASSIFICATION SCHEDULE
This policy establishes mandatory retention schedules for corporate records in accordance with statutory, regulatory, and operational requirements. Records are classified into the following retention periods:
- Financial & Tax Documents: 7 years from tax filing date (audits, ledgers, tax returns).
- Executed Contracts & Agreements: 10 years following expiration or formal termination.
- Employee & Personnel Records: 7 years following date of employee separation.
- Technical Logs & Audit Trails: 1 year in active queryable storage; 3 years in cold archival.
- Corporate Governance & Board Minutes: Permanent retention in corporate legal archives.

## SECTION 2: LITIGATION HOLD PROCEDURES
Upon reasonable anticipation of litigation, regulatory subpoena, or government investigation, the General Counsel issues an immediate 'Litigation Hold' notice. All routine disposal and automated data deletion routines for relevant documents are immediately suspended until the hold is formally rescinded in writing.

## SECTION 3: SECURE DESTRUCTION STANDARDS
When records reach the end of their retention lifecycle, electronic media must be sanitized in compliance with NIST SP 800-88 Guidelines for Media Sanitization (cryptographic erasure or physical destruction)."""
        },
        {
            "filename": "Anti_Bribery_and_Corruption_Policy.pdf",
            "type": "pdf",
            "title": "Nexus Technologies - Global Anti-Bribery & Anti-Corruption Policy",
            "text": """# Nexus Technologies - Global Anti-Bribery & Anti-Corruption Policy

## SECTION 1: POLICY STATEMENT & STATUTORY COMPLIANCE
Nexus enforces zero tolerance for bribery, extortion, kickbacks, or corrupt practices in any business dealing. Employees, directors, and representatives must strictly comply with the US Foreign Corrupt Practices Act (FCPA), the UK Bribery Act 2010, and all applicable anti-corruption legislation in territories where Nexus operates.

## SECTION 2: PROHIBITION OF BRIBES AND FACILITATION PAYMENTS
Employees are prohibited from directly or indirectly giving, offering, promising, or authorizing anything of value (including money, gifts, hospitality, loans, or employment opportunities) to any government official or private counterparty to influence an official act or secure an improper business advantage. Facilitation or 'grease' payments to expedite routine government actions are strictly illegal and prohibited.

## SECTION 3: GIFTS, MEALS, AND ENTERTAINMENT THRESHOLDS
Gifts and entertainment must be customary, reasonable, and modest:
- Maximum Gift Limit: $100.00 USD per individual per calendar year.
- Any gift or business meal involving a government official requires prior written pre-clearance from the Legal Compliance Officer.
- Cash or cash equivalents (gift cards, vouchers) are strictly prohibited under all circumstances.

## SECTION 4: DUE DILIGENCE ON THIRD PARTIES AND INTERMEDIARIES
Third-party agents, channel partners, and distributors must undergo comprehensive anti-corruption background checks and execute contractual anti-bribery covenants prior to engagement."""
        }
    ]
}


def main():
    logger.info("================================================================================")
    logger.info("STARTING FAKE COMPANY DATA GENERATION & SEEDING FOR NEXUS TECHNOLOGIES INC.")
    logger.info("================================================================================")

    # Base folder for documents
    docs_base_dir = os.path.join(PROJECT_ROOT, "data", "company_documents")
    os.makedirs(docs_base_dir, exist_ok=True)

    # Initialize Qdrant collection
    logger.info("Initializing Qdrant collection 'hybrid_knowledge_base'...")
    try:
        init_qdrant()
        logger.info("Qdrant collection ready.")
    except Exception as e:
        logger.error(f"Error initializing Qdrant: {e}")

    total_files_created = 0
    total_chunks_indexed = 0

    for dept, file_list in DOCUMENTS.items():
        dept_dir = os.path.join(docs_base_dir, dept)
        os.makedirs(dept_dir, exist_ok=True)
        logger.info(f"\n---> Processing department: '{dept}' ({len(file_list)} documents)")

        for doc_info in file_list:
            fname = doc_info["filename"]
            ftype = doc_info["type"]
            fpath = os.path.join(dept_dir, fname)
            title = doc_info.get("title", fname)
            content_text = doc_info.get("text", "")

            # 1. Create Physical File
            if ftype == "pdf":
                paragraphs = [line for line in content_text.split('\n') if line.strip()]
                build_pdf(fpath, title, paragraphs)
            elif ftype == "docx":
                sections = doc_info.get("sections", [("Document Content", content_text)])
                build_docx(fpath, title, sections)
            elif ftype == "xlsx":
                sheets_data = doc_info.get("sheets", {"Sheet1": [["Data", "Value"], ["Sample", 100]]})
                build_xlsx(fpath, sheets_data)
            elif ftype == "md":
                build_md(fpath, content_text)

            logger.info(f"  [CREATED] {dept}/{fname} ({os.path.getsize(fpath)} bytes)")
            total_files_created += 1

            # 2. Chunk text and upsert directly into Qdrant
            chunks = chunk_document(content_text, fname)
            try:
                upsert_document_chunks(chunks, fname, department=dept)
                total_chunks_indexed += len(chunks)
                logger.info(f"  [INDEXED] {fname} -> {len(chunks)} chunks into Qdrant (dept='{dept}')")
            except Exception as e:
                logger.error(f"  [FAILED TO INDEX] {fname}: {e}")

    logger.info("================================================================================")
    logger.info(f"SEEDING COMPLETE! Created {total_files_created} files, indexed {total_chunks_indexed} chunks.")
    logger.info("================================================================================")


if __name__ == "__main__":
    main()
